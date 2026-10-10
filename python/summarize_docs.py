"""批量总结一个文件夹里的 PDF / Word / txt / md 文档，结果写进 summary.csv（Excel 可直接打开）。

用法：
    pip install -U openai python-docx
    python summarize_docs.py 文件夹路径
    python summarize_docs.py 文件夹路径 --model gpt-6.1-sol --question "这份材料对我们的产品有什么启发？"
"""
import argparse, base64, csv, json, os, sys, time
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from openai import OpenAI

client = OpenAI(base_url="https://hivegpt.cn/v1", api_key=os.environ["HIVEGPT_API_KEY"], timeout=300)

WORKERS = 3              # 同时处理几个文件；遇到 429 就调小
MAX_CHARS = 60_000       # Word / 文本只取前这么多字，防止超长文档花太多钱
SCHEMA = {
    "type": "object",
    "properties": {
        "title": {"type": "string", "description": "文档标题，没有就根据内容起一个"},
        "summary": {"type": "string", "description": "一句话摘要，50 字以内"},
        "key_points": {"type": "array", "items": {"type": "string"}, "description": "3 到 5 条要点"},
        "keywords": {"type": "array", "items": {"type": "string"}},
        "answer": {"type": "string", "description": "对用户问题的回答；没有问题就留空"},
    },
    "required": ["title", "summary", "key_points", "keywords", "answer"],
    "additionalProperties": False,
}


def file_content(path: Path):
    """把文件变成消息里的一段内容：PDF 整个发过去，其他格式取文字。"""
    ext = path.suffix.lower()
    if ext == ".pdf":
        data = base64.b64encode(path.read_bytes()).decode()
        return {"type": "file", "file": {"filename": path.name, "file_data": "data:application/pdf;base64," + data}}
    if ext == ".docx":
        import docx  # pip install python-docx
        doc = docx.Document(str(path))
        text = "\n".join(p.text for p in doc.paragraphs)
        for table in doc.tables:  # 表格里的文字也带上
            for row in table.rows:
                text += "\n" + " | ".join(cell.text for cell in row.cells)
    else:
        text = path.read_text(encoding="utf-8", errors="ignore")
    return {"type": "text", "text": f"文件名：{path.name}\n\n{text[:MAX_CHARS]}"}


def summarize(path: Path, model: str, question: str):
    prompt = "阅读这份文档，按要求的格式输出中文总结。"
    if question:
        prompt += f"\n另外回答这个问题（写在 answer 里）：{question}"
    for attempt in range(3):
        try:
            resp = client.chat.completions.create(
                model=model,
                messages=[{"role": "user", "content": [file_content(path), {"type": "text", "text": prompt}]}],
                response_format={"type": "json_schema", "json_schema": {"name": "doc", "strict": True, "schema": SCHEMA}},
                reasoning_effort="low",
            )
            r = json.loads(resp.choices[0].message.content)
            r["file"], r["tokens"] = path.name, resp.usage.total_tokens
            print(f"✓ {path.name}（{resp.usage.total_tokens} token）", file=sys.stderr)
            return r
        except Exception as e:
            print(f"  {path.name} 出错：{e}，重试", file=sys.stderr)
            time.sleep(2 ** attempt)
    return {"file": path.name, "title": "", "summary": "处理失败", "key_points": [], "keywords": [], "answer": "", "tokens": 0}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("folder")
    ap.add_argument("--model", default="gpt-6-luna")
    ap.add_argument("--question", default="")
    args = ap.parse_args()

    files = sorted(p for p in Path(args.folder).iterdir() if p.suffix.lower() in {".pdf", ".docx", ".txt", ".md"})
    print(f"共 {len(files)} 个文件", file=sys.stderr)
    with ThreadPoolExecutor(WORKERS) as pool:
        results = list(pool.map(lambda p: summarize(p, args.model, args.question), files))

    with open("summary.csv", "w", newline="", encoding="utf-8-sig") as f:  # utf-8-sig：Excel 打开不乱码
        w = csv.writer(f)
        w.writerow(["文件", "标题", "一句话摘要", "要点", "关键词", "问题回答", "token"])
        for r in results:
            w.writerow([r["file"], r["title"], r["summary"], "\n".join(r["key_points"]), "、".join(r["keywords"]), r["answer"], r["tokens"]])
    print(f"完成：summary.csv，共用 {sum(r['tokens'] for r in results)} token", file=sys.stderr)


if __name__ == "__main__":
    main()
